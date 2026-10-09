"""Conversion au format pérenne PDF/A + OCR + extraction du texte pour l'indexation.

Chaîne de traitement (dégradation gracieuse si un outil manque) :
  * Bureautique (docx, odt, xlsx, pptx, rtf...) -> LibreOffice headless, filtre PDF/A-2b
  * PDF                                          -> OCRmyPDF --output-type pdfa-2 --skip-text
  * Images / scans (png, jpg, tiff)              -> OCRmyPDF (image -> PDF/A-2b avec couche texte)
                                                    à défaut Tesseract -> PDF recherchable
  * Texte extrait : pypdf sur le PDF produit ; si vide -> OCR Tesseract page par page.
"""
import mimetypes
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from app.core.config import get_settings

OFFICE_EXT = {".doc", ".docx", ".odt", ".rtf", ".xls", ".xlsx", ".ods", ".ppt", ".pptx", ".odp"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_EXT = {".txt", ".csv", ".md", ".xml", ".json", ".html", ".htm"}
MAX_TEXT = 2_000_000  # caractères indexés par version


@dataclass
class ConversionResult:
    pdf_path: Path | None = None
    archival_format: str | None = None
    text: str = ""
    ocr_applied: bool = False
    status: str = "SKIPPED"
    log: list[str] = field(default_factory=list)


def guess_mime(filename: str, declared: str | None, head: bytes) -> str:
    if head.startswith(b"%PDF"):
        return "application/pdf"
    if head.startswith(b"\x89PNG"):
        return "image/png"
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "image/tiff"
    guessed = mimetypes.guess_type(filename)[0]
    if guessed:
        return guessed
    if declared and declared != "application/octet-stream":
        return declared
    return "application/octet-stream"


def tool(name: str) -> str | None:
    return shutil.which(name)


def _run(cmd: list[str], timeout: int = 600) -> tuple[bool, str]:
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        return p.returncode == 0, (p.stderr or p.stdout)[-800:]
    except (subprocess.TimeoutExpired, OSError) as exc:
        return False, str(exc)


def _ocrmypdf(src: Path, out: Path, langs: str, res: ConversionResult, image: bool = False) -> bool:
    if not tool("ocrmypdf"):
        return False
    cmd = ["ocrmypdf", "--output-type", "pdfa-2", "-l", langs, "--skip-text", "--optimize", "1"]
    if image:
        cmd += ["--image-dpi", "300"]
    ok, msg = _run(cmd + [str(src), str(out)])
    if not ok:
        # Sans Ghostscript, OCRmyPDF ne sait pas produire de PDF/A : on garde un PDF OCRisé
        cmd[2] = "pdf"
        ok2, msg2 = _run(cmd + [str(src), str(out)])
        res.log.append(f"ocrmypdf pdfa échec ({msg.strip()[:200]}) ; repli pdf : {'ok' if ok2 else msg2[:200]}")
        if ok2:
            res.archival_format = "PDF (OCR)"
            res.ocr_applied = True
        return ok2
    res.archival_format = "PDF/A-2b"
    res.ocr_applied = True
    res.log.append("ocrmypdf : PDF/A-2b produit")
    return True


def _libreoffice(src: Path, outdir: Path, res: ConversionResult) -> Path | None:
    exe = tool("soffice") or tool("libreoffice")
    if not exe:
        res.log.append("LibreOffice absent : pas de conversion bureautique")
        return None
    filt = 'pdf:writer_pdf_Export:{"SelectPdfVersion":{"type":"long","value":"2"}}'
    ok, msg = _run([exe, "--headless", "--norestore", f"-env:UserInstallation=file://{outdir}/lo-profile",
                    "--convert-to", filt, "--outdir", str(outdir), str(src)], timeout=180)
    out = outdir / (src.stem + ".pdf")
    if ok and out.exists():
        res.archival_format = "PDF/A-2b"
        res.log.append("LibreOffice : PDF/A-2b produit")
        return out
    res.log.append(f"LibreOffice échec : {msg[:200]}")
    return None


def _tesseract_image_to_pdf(src: Path, outbase: Path, langs: str, res: ConversionResult) -> Path | None:
    if not tool("tesseract"):
        return None
    ok, msg = _run(["tesseract", str(src), str(outbase), "-l", _available_langs(langs), "pdf"])
    out = outbase.with_suffix(".pdf")
    if ok and out.exists():
        res.archival_format, res.ocr_applied = "PDF (OCR)", True
        res.log.append("tesseract : PDF recherchable produit (PDF/A nécessite OCRmyPDF + Ghostscript)")
        return out
    res.log.append(f"tesseract échec : {msg[:200]}")
    return None


def _available_langs(langs: str) -> str:
    try:
        import pytesseract
        have = set(pytesseract.get_languages(config=""))
        keep = [lang for lang in langs.split("+") if lang in have]
        return "+".join(keep) or "eng"
    except Exception:
        return langs


def extract_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
            if sum(map(len, parts)) > MAX_TEXT:
                break
        return "\n".join(parts).strip()
    except Exception:
        return ""


def ocr_pdf_pages(path: Path, langs: str, max_pages: int = 50) -> str:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:
        return ""
    pdf = pdfium.PdfDocument(str(path))
    out = []
    for i in range(min(len(pdf), max_pages)):
        img = pdf[i].render(scale=300 / 72).to_pil()
        out.append(pytesseract.image_to_string(img, lang=_available_langs(langs)))
    return "\n".join(out).strip()


def ocr_image(path: Path, langs: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path), lang=_available_langs(langs)).strip()
    except Exception:
        return ""


def extract_office_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        try:
            import docx
            d = docx.Document(str(path))
            return "\n".join(p.text for p in d.paragraphs).strip()
        except Exception:
            return ""
    return ""


def process(src: Path, filename: str, mime: str, workdir: Path, convert: bool | None = None) -> ConversionResult:
    """Point d'entrée : produit la représentation PDF/A (si possible) et le texte indexable."""
    s = get_settings()
    convert = s.enable_conversion if convert is None else convert
    langs = s.ocr_languages
    ext = Path(filename).suffix.lower()
    res = ConversionResult()
    # Copie avec la bonne extension (requis par LibreOffice/Tesseract)
    work_src = workdir / ("source" + (ext or ".bin"))
    shutil.copyfile(src, work_src)

    if ext in TEXT_EXT or mime.startswith("text/"):
        res.text = work_src.read_text(errors="replace")[:MAX_TEXT]
        if convert:
            pdf = _libreoffice(work_src, workdir, res)
            res.pdf_path = pdf
        res.status = "DONE" if res.pdf_path else "PARTIAL"
        return res

    if mime == "application/pdf":
        if convert:
            out = workdir / "archival.pdf"
            if _ocrmypdf(work_src, out, langs, res):
                res.pdf_path = out
        target = res.pdf_path or work_src
        res.text = extract_pdf_text(target)
        if len(res.text) < 20:  # PDF image (scan) sans couche texte
            ocr = ocr_pdf_pages(work_src, langs)
            if ocr:
                res.text, res.ocr_applied = ocr, True
                res.log.append("OCR tesseract page par page")
    elif mime.startswith("image/") or ext in IMAGE_EXT:
        if convert:
            out = workdir / "archival.pdf"
            if _ocrmypdf(work_src, out, langs, res, image=True):
                res.pdf_path = out
            else:
                res.pdf_path = _tesseract_image_to_pdf(work_src, workdir / "archival", langs, res)
        res.text = extract_pdf_text(res.pdf_path) if res.pdf_path else ""
        if not res.text:
            res.text = ocr_image(work_src, langs)
            res.ocr_applied = res.ocr_applied or bool(res.text)
    elif ext in OFFICE_EXT:
        if convert:
            res.pdf_path = _libreoffice(work_src, workdir, res)
        res.text = extract_pdf_text(res.pdf_path) if res.pdf_path else extract_office_text(work_src)
    else:
        res.log.append(f"Format non pris en charge pour la conversion : {mime}")
        res.status = "SKIPPED"
        return res

    res.text = (res.text or "")[:MAX_TEXT]
    if res.pdf_path and res.archival_format == "PDF/A-2b":
        res.status = "DONE"
    elif res.pdf_path or res.text:
        res.status = "PARTIAL"
    else:
        res.status = "FAILED" if convert else "SKIPPED"
    return res


def new_workdir() -> tempfile.TemporaryDirectory:
    return tempfile.TemporaryDirectory(prefix="siged-")
